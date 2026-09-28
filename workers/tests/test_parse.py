from __future__ import annotations

import uuid
from pathlib import Path

import pytest
from docx import Document

from inspector_workers.contracts import validate
from inspector_workers.parse.docx import parse_docx
from inspector_workers.parse.layout import norm_bbox
from inspector_workers.parse.pdf import parse_pdf
from inspector_workers.parse.worker import ParseWorker, detect_kind
from inspector_workers.parse.xmlfile import parse_xml

TEP = ["Технико-экономические показатели", "Площадь застройки — 1 250,5 м²", "Общая площадь здания: 8 400 м²",
       "Количество этажей: 9"]


def test_norm_bbox_is_valid_for_degenerate_and_out_of_range():
    for box in (norm_bbox(10, 10, 10, 10, 100, 100), norm_bbox(-5, -5, 300, 300, 100, 100), norm_bbox(100, 100, 100, 100, 100, 100)):
        x0, y0, x1, y1 = box
        assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1


def test_parse_pdf_lines_and_bbox(make_pdf):
    path = make_pdf([TEP, ["Второй лист", "Ширина двери 0,8 м на первом этаже здания"]])
    res = parse_pdf(path, None, None, use_ocr=False)
    assert res.quality == "OK" and len(res.pages) == 2
    texts = [ln.text for ln in res.pages[0].lines]
    assert "Площадь застройки — 1 250,5 м²" in texts
    for page in res.pages:
        ys = [ln.bbox[1] for ln in page.lines]
        assert ys == sorted(ys), "строки идут сверху вниз"
        for ln in page.lines:
            x0, y0, x1, y1 = ln.bbox
            assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1


def test_parse_pdf_page_range(make_pdf):
    path = make_pdf([[f"Страница номер {i} " + "текст " * 8] for i in range(1, 6)])
    res = parse_pdf(path, 2, 3, use_ocr=False)
    assert [p.page for p in res.pages] == [2, 3], "номера страниц — абсолютные, не относительные"
    res = parse_pdf(path, 4, 99, use_ocr=False)
    assert [p.page for p in res.pages] == [4, 5], "page_to за пределами клампится"


def test_scan_pages_downgrade_quality(make_pdf):
    path = make_pdf([TEP, []])
    res = parse_pdf(path, None, None, use_ocr=False)
    assert res.quality == "LOW_QUALITY"
    only_scan = parse_pdf(make_pdf([[], []]), None, None, use_ocr=False)
    assert only_scan.quality == "ABSTAIN"


def test_parse_docx_paragraphs_and_tables(tmp_path):
    doc = Document()
    doc.add_paragraph("Пояснительная записка")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Общая площадь здания", "8 400 м²"
    t.cell(1, 0).text, t.cell(1, 1).text = "Этажность", "9"
    p = tmp_path / "a.docx"
    doc.save(p)
    res = parse_docx(p)
    assert res.quality == "OK"
    texts = [ln.text for ln in res.pages[0].lines]
    assert texts == ["Пояснительная записка", "Общая площадь здания | 8 400 м²", "Этажность | 9"]
    assert all(ln.bbox is None for ln in res.pages[0].lines)


def test_parse_xml(tmp_path):
    p = tmp_path / "a.xml"
    p.write_text('<?xml version="1.0" encoding="utf-8"?><Doc xmlns="urn:x"><Building><Area>1250,5</Area>'
                 '<Name>Корпус 1 жилой дом</Name></Building></Doc>', encoding="utf-8")
    res = parse_xml(p)
    assert [ln.text for ln in res.pages[0].lines] == ["Doc/Building/Area: 1250,5",
                                                      "Doc/Building/Name: Корпус 1 жилой дом"]


def test_detect_kind_by_content_not_extension(make_pdf, tmp_path):
    pdf = make_pdf([TEP])
    renamed = tmp_path / "scan.docx"
    renamed.write_bytes(pdf.read_bytes())
    assert detect_kind(renamed) == "pdf"
    junk = tmp_path / "x.pdf"
    junk.write_bytes(b"MZ\x90\x00garbage")
    with pytest.raises(ValueError):
        detect_kind(junk)


def _env(data: dict) -> dict:
    return {"event_id": str(uuid.uuid4()), "event_type": "doc.parse.requested", "schema_version": 1,
            "occurred_at": "2026-09-28T10:00:00Z", "correlation_id": "p", "object_id": str(uuid.uuid4()),
            "process_id": str(uuid.uuid4()), "data": data}


def test_worker_produces_contract_valid_event_and_layout(settings, storage, make_pdf):
    pdf = make_pdf([TEP])
    storage.files["k/a.pdf"] = pdf.read_bytes()
    job, file_id = str(uuid.uuid4()), str(uuid.uuid4())
    env = _env({"job_id": job, "file_id": file_id, "doc_stage": "PD", "discipline": "AR", "sha256": "a" * 64,
                "storage_key": "k/a.pdf", "content_type": "application/pdf", "page_from": 1, "page_to": 20})
    validate(settings.contracts_dir, "doc.parse.requested", env)

    outs = ParseWorker(settings, storage).process(env)
    assert len(outs) == 1 and outs[0].topic == "doc.parse.completed"
    data = outs[0].data
    assert data["quality"] == "OK" and data["layout_key"] == f"layouts/{job}.json"

    completed = {**env, "event_type": "doc.parse.completed", "data": data}
    validate(settings.contracts_dir, "doc.parse.completed", completed)

    layout = storage.json[data["layout_key"]]
    assert layout["sha256"] == "a" * 64 and layout["stage"] == "PD" and layout["file_id"] == file_id
    assert layout["pages"][0]["lines"][1]["text"] == "Площадь застройки — 1 250,5 м²"


def test_worker_failure_event_is_terminal_and_valid(settings, storage):
    env = _env({"job_id": str(uuid.uuid4()), "file_id": str(uuid.uuid4()), "storage_key": "k/x", "page_from": None, "page_to": None})
    outs = ParseWorker(settings, storage).on_failure(env, RuntimeError("boom"))
    assert outs[0].data["quality"] == "FAILED" and outs[0].data["layout_key"] is None
    validate(settings.contracts_dir, "doc.parse.completed", {**env, "event_type": "doc.parse.completed", "data": outs[0].data})


def test_short_docx_and_xml_are_not_abstained(tmp_path):
    d = Document(); d.add_paragraph("Коротко")
    p = tmp_path / "s.docx"; d.save(p)
    assert parse_docx(p).quality == "OK"
    x = tmp_path / "s.xml"; x.write_text("<a><b>1</b></a>", encoding="utf-8")
    assert parse_xml(x).quality == "OK"
    empty = Document(); e = tmp_path / "e.docx"; empty.save(e)
    assert parse_docx(e).quality == "ABSTAIN", "пустой документ — по-прежнему ABSTAIN"
